"""Report export — BOQ + financial summary as xlsx."""

from __future__ import annotations

import io
import uuid

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from app.models.cost import CostLine
from app.models.project import Project
from app.models.takeoff import TakeoffItem
from app.services.cost_engine import recalculate_project_cost


def build_boq_xlsx(db: Session, *, organization_id: uuid.UUID, project_id: uuid.UUID) -> bytes:
    project = db.get(Project, project_id)
    if project is None:
        raise ValueError("project not found")

    summary = recalculate_project_cost(
        db, organization_id=organization_id, project_id=project_id
    )

    wb = Workbook()

    # ---------- BOQ sheet ----------
    ws = wb.active
    assert ws is not None
    ws.title = "BOQ"
    ws.sheet_view.rightToLeft = True

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    headers = [
        "#",
        "الفئة",
        "الوصف",
        "الوحدة",
        "الكمية",
        "الثقة",
        "حالة المراجعة",
        "تكلفة المواد",
        "تكلفة العمالة",
        "تكلفة المعدات",
        "إجمالي البند",
    ]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")

    rows = (
        db.query(TakeoffItem, CostLine)
        .outerjoin(CostLine, CostLine.takeoff_item_id == TakeoffItem.id)
        .filter(TakeoffItem.project_id == project_id)
        .all()
    )
    for idx, (item, cl) in enumerate(rows, start=1):
        ws.cell(row=idx + 1, column=1, value=idx)
        ws.cell(row=idx + 1, column=2, value=item.category)
        ws.cell(row=idx + 1, column=3, value=item.description or "")
        ws.cell(row=idx + 1, column=4, value=item.unit)
        ws.cell(row=idx + 1, column=5, value=float(item.quantity))
        ws.cell(row=idx + 1, column=6, value=float(item.confidence))
        ws.cell(row=idx + 1, column=7, value=item.review_status)
        ws.cell(row=idx + 1, column=8, value=float(cl.material_cost) if cl else None)
        ws.cell(row=idx + 1, column=9, value=float(cl.labor_cost) if cl else None)
        ws.cell(row=idx + 1, column=10, value=float(cl.equipment_cost) if cl else None)
        ws.cell(row=idx + 1, column=11, value=float(cl.line_total) if cl else None)

    for col_letter in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K"]:
        ws.column_dimensions[col_letter].width = 16

    # ---------- Summary sheet ----------
    s = wb.create_sheet("الملخص المالي")
    s.sheet_view.rightToLeft = True
    s["A1"] = "المشروع"
    s["B1"] = project.name
    s["A2"] = "العميل"
    s["B2"] = project.client_name or "-"
    s["A3"] = "المدينة"
    s["B3"] = project.location_city or "-"
    s["A4"] = "المساحة الإجمالية (م²)"
    s["B4"] = float(project.gross_area_m2) if project.gross_area_m2 else None

    s["A6"] = "التكلفة المباشرة"
    s["B6"] = float(summary.direct)
    s["A7"] = "غير المباشرة"
    s["B7"] = float(summary.indirect)
    s["A8"] = "الطوارئ"
    s["B8"] = float(summary.contingency)
    s["A9"] = "الإجمالي"
    s["B9"] = float(summary.total)
    s["A10"] = "هامش الربح"
    s["B10"] = float(summary.margin_amount)
    s["A11"] = "سعر العرض"
    s["B11"] = float(summary.bid_price)
    s["A12"] = "تكلفة م²"
    s["B12"] = float(summary.cost_per_m2) if summary.cost_per_m2 is not None else None

    for r in range(6, 13):
        s.cell(row=r, column=1).font = Font(bold=True)
    s.column_dimensions["A"].width = 24
    s.column_dimensions["B"].width = 20

    # ---------- By category ----------
    c = wb.create_sheet("حسب الفئة")
    c.sheet_view.rightToLeft = True
    c["A1"] = "الفئة"
    c["B1"] = "التكلفة"
    c["A1"].font = Font(bold=True)
    c["B1"].font = Font(bold=True)
    for i, (cat, val) in enumerate(sorted(summary.by_category.items()), start=2):
        c.cell(row=i, column=1, value=cat)
        c.cell(row=i, column=2, value=float(val))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
